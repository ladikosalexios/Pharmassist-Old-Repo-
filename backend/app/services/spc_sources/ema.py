"""EMA adapter — Greek product-information PDFs for centrally-authorised drugs.

Resolution path: EMA publishes a downloadable medicines dataset (xlsx, URL in
``settings.spc_ema_dataset_url`` because it drifts); we cache it in-memory for
24h as {normalised name → medicine page URL}, match the catalog row's INN
(``name_en``) or Greek brand prefix, then derive the combined SmPC+PIL Greek
product-information PDF (``…-epar-product-information_el.pdf``), HEAD-checking
and falling back to scraping the medicine page for the ``_el.pdf`` link.

EMA documents are COMBINED (SmPC + annexes + PIL in one file) → one FoundDoc
with ``doc_type="combined"``; both splitters run over it at ingest. All
failures degrade to AdapterResult.error — never an exception.
"""

import io
import logging
import re
import time

from openpyxl import load_workbook

from ...config import get_settings
from ...db.models.drug_catalog import DrugCatalog
from .base import AdapterResult, FoundDoc, _client, _throttle

logger = logging.getLogger(__name__)

SOURCE = "ema"
_DATASET_TTL_SECONDS = 24 * 3600

# {normalised medicine name → medicine page URL}; module-level 24h cache.
_dataset: dict[str, str] | None = None
_dataset_loaded_at: float = 0.0

_EL_PDF_RE = re.compile(r'href="([^"]+product-information_el\.pdf[^"]*)"')


def _norm(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.lower())


def _slug(medicine_url: str) -> str | None:
    m = re.search(r"/medicines/human/(?:EPAR/)?([a-z0-9-]+)", medicine_url, re.IGNORECASE)
    return m.group(1) if m else None


async def _load_dataset() -> dict[str, str]:
    """The EMA medicines dataset as {normalised name → medicine page URL}."""
    global _dataset, _dataset_loaded_at
    if _dataset is not None and time.monotonic() - _dataset_loaded_at < _DATASET_TTL_SECONDS:
        return _dataset

    settings = get_settings()
    await _throttle(SOURCE)
    async with _client(follow_redirects=True) as client:
        r = await client.get(settings.spc_ema_dataset_url)
        r.raise_for_status()
        content = r.content

    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    mapping: dict[str, str] = {}
    name_col = url_col = None
    for row in ws.iter_rows(values_only=True):
        cells = [str(c) if c is not None else "" for c in row]
        if name_col is None:
            # Locate the header row by its column names (position drifts).
            lowered = [c.lower() for c in cells]
            for i, c in enumerate(lowered):
                if "medicine" in c and "name" in c:
                    name_col = i
                if c.startswith("url") or "medicine url" in c:
                    url_col = i
            continue
        if url_col is None or name_col >= len(cells) or url_col >= len(cells):
            continue
        name, url = cells[name_col], cells[url_col]
        if name and url.startswith("http"):
            mapping[_norm(name)] = url
    wb.close()

    _dataset = mapping
    _dataset_loaded_at = time.monotonic()
    logger.info("[spc][ema] dataset loaded: %d medicines", len(mapping))
    return mapping


def _match(dataset: dict[str, str], product: DrugCatalog) -> str | None:
    for candidate in (product.name_en, (product.name_gr or "").split(" ")[0]):
        if not candidate:
            continue
        key = _norm(candidate)
        if not key:
            continue
        if key in dataset:
            return dataset[key]
        for name, url in dataset.items():
            if name.startswith(key) or key.startswith(name):
                return url
    return None


async def resolve(product: DrugCatalog) -> AdapterResult:
    """Find the Greek product-information PDF for one product. Never raises."""
    try:
        dataset = await _load_dataset()
        page_url = _match(dataset, product)
        if page_url is None:
            return AdapterResult(error="ema: no dataset match")
        slug = _slug(page_url)

        async with _client(follow_redirects=True) as client:
            if slug:
                pdf_url = (
                    "https://www.ema.europa.eu/en/documents/product-information/"
                    f"{slug}-epar-product-information_el.pdf"
                )
                await _throttle(SOURCE)
                head = await client.head(pdf_url)
                if head.status_code == 200:
                    return AdapterResult(
                        docs=[FoundDoc(url=pdf_url, doc_type="combined", source=SOURCE)]
                    )
            # Fallback: scrape the medicine page for the _el.pdf link.
            await _throttle(SOURCE)
            page = await client.get(page_url)
            page.raise_for_status()
            m = _EL_PDF_RE.search(page.text)
            if not m:
                return AdapterResult(error="ema: no Greek product-information PDF on page")
            href = m.group(1)
            if href.startswith("/"):
                href = f"https://www.ema.europa.eu{href}"
            return AdapterResult(docs=[FoundDoc(url=href, doc_type="combined", source=SOURCE)])
    except Exception as exc:
        logger.warning("[spc][ema] resolve failed for %s: %s", product.gns_code, exc)
        return AdapterResult(error=f"ema: {type(exc).__name__}: {exc}")
